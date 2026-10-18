"""Labeler ki xlsx sheet ko data/test_fresh.jsonl mein badalta hai.

Use:
  python scripts/make_fresh.py labelled.xlsx answer_key.csv data/test_fresh.jsonl
"""
import sys, json, csv
from collections import Counter
from openpyxl import load_workbook

LABELS = {"scam", "grooming", "bullying", "inappropriate", "safe"}
# ASSUMPTION: labeler risk nahi deta, isliye label se risk set kiya hai.
# Apne validate.py ke allowed values se match karo, warna yahan badlo.
RISK = {"safe": "low", "bullying": "medium", "inappropriate": "medium",
        "scam": "high", "grooming": "high"}

sheet_path, key_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]

ws = load_workbook(sheet_path)["Messages"]
head = [c.value for c in ws[1]]
col = {h: i for i, h in enumerate(head)}

key = {}
with open(key_path, encoding="utf-8-sig", newline="") as f:
    for r in csv.DictReader(f):
        key[r["id"]] = r["my_label"].strip().lower()

out, skipped, disagree = [], [], []
for row in ws.iter_rows(min_row=2, values_only=True):
    rid, text = row[col["id"]], row[col["message"]]
    if not rid or not text:
        continue
    label = (row[col["label"]] or "").strip().lower()
    if label not in LABELS:
        skipped.append(rid)
        continue
    unsure = (row[col["unsure"]] or "").strip().lower() == "yes"
    note = (row[col["labeler_note"]] or "").strip()
    lang = (row[col["language"]] or "").strip().lower()
    rec = {
        "id": rid,
        "text": text,
        "language": lang,
        "label": label,
        "risk": RISK[label],
        "reason": note or "labeler ne diya",
        "source": "test_fresh",
        "unsure": unsure,
    }
    out.append(rec)
    if key.get(rid) and key[rid] != label:
        disagree.append((rid, key[rid], label, unsure, text))

with open(out_path, "w", encoding="utf-8") as f:
    for r in out:
        f.write(json.dumps(r, ensure_ascii=False) + "\n")

print(f"wrote {len(out)} rows to {out_path}")
print("label counts:", dict(Counter(r["label"] for r in out)))
print("unsure rows:", sum(r["unsure"] for r in out))
print("skipped (label khaali ya galat):", skipped)
print(f"\nMere label se alag (labeler ka label rakha gaya): {len(disagree)}")
for rid, mine, theirs, unsure, text in disagree:
    print(f"- {rid}: mera={mine} labeler={theirs} unsure={unsure}")
    print(f"    {text}")